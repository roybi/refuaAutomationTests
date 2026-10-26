import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest.mock import patch

import yaml
from defusedxml import ElementTree
from openpyxl import Workbook, load_workbook

import azure_testcase_importer as importer
import local_upload
import workflow_runner

WORKBOOK = Path(__file__).with_name("Meditik_Test_Cases.xlsx")
FIXTURE_DIRECTORY = None


def setUpModule():
    global WORKBOOK, FIXTURE_DIRECTORY
    FIXTURE_DIRECTORY = tempfile.TemporaryDirectory()
    WORKBOOK = Path(FIXTURE_DIRECTORY.name, "test_cases.xlsx")
    workbook = Workbook()
    sheet = workbook.create_sheet("Test_Cases")
    sheet.append(["TC_ID", "Scenario", "Feature", "Preconditions", "Priority", "Tags",
                  "Test Data", "Detailed Automation Steps", "Detailed Expected Results"])
    bdd = workbook.create_sheet("BDD_Test_Cases")
    bdd.append(["TC_ID", "Scenario", "Feature", "Pre-Steps", "When", "Then"])
    for index in range(1, 15):
        test_id = f"SAMPLE-{index:03d}"
        sheet.append([test_id, f"Scenario {index}", "Medical Profile", "Signed in", "P1",
                      "smoke", "Synthetic data", "1. Open page\n2. Submit", "Success"])
        bdd.append([test_id, f"Scenario {index}", "Medical Profile", "Signed in", "Submit", "Success"])
    workbook.save(WORKBOOK)
    workbook.close()


def tearDownModule():
    if FIXTURE_DIRECTORY is not None:
        FIXTURE_DIRECTORY.cleanup()


class PlanNameTests(unittest.TestCase):
    def test_prefix_is_added_once(self):
        for name in ("Visits", "AUTO_Visits", "auto_Visits", " Auto_AUTO_Visits "):
            with self.subTest(name=name):
                self.assertEqual(importer.normalize_plan_name(name), "AUTO_Visits")

    def test_empty_name_is_rejected(self):
        for name in ("", "  ", "AUTO_", "AUTO_AUTO_"):
            with self.subTest(name=name), self.assertRaises(ValueError):
                importer.normalize_plan_name(name)

    def test_existing_mixed_case_prefix_is_reused(self):
        client = importer.AzureDevOpsClient("example", "example", "test-token")
        self.addCleanup(client.session.close)
        existing = {"id": 100, "name": "Auto_Visits"}
        with patch.object(client, "request", return_value={"value": [existing]}) as request:
            with redirect_stdout(io.StringIO()):
                self.assertEqual(client.get_or_create_plan("Visits"), existing)
        request.assert_called_once()

    def test_new_plan_has_uppercase_prefix(self):
        client = importer.AzureDevOpsClient("example", "example", "test-token")
        self.addCleanup(client.session.close)
        with patch.object(client, "request", side_effect=[{"value": []}, {"id": 100}]) as request:
            with redirect_stdout(io.StringIO()):
                client.get_or_create_plan("Visits")
        self.assertEqual(request.call_args.kwargs["json_body"]["name"], "AUTO_Visits")


class ExcelImportTests(unittest.TestCase):
    def test_workbook_preserves_detailed_columns(self):
        cases = importer.read_excel(WORKBOOK, "Test Cases", None, False)
        workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
        try:
            rows = workbook["Test_Cases"].iter_rows(values_only=True)
            headers = next(rows)
            records = [dict(zip(headers, row)) for row in rows]
        finally:
            workbook.close()
        records = [record for record in records if record["Scenario"]]
        self.assertEqual(len(cases), 14)
        self.assertEqual(len(cases), len(records))
        for case, record in zip(cases, records):
            with self.subTest(test_id=case.test_id):
                self.assertEqual(case.test_id, record["TC_ID"])
                self.assertEqual(case.title, record["Scenario"])
                self.assertEqual(case.module, record["Feature"])
                self.assertEqual(case.actions, importer.split_multiline(record["Detailed Automation Steps"]))
                self.assertEqual(case.expected, importer.split_multiline(record["Detailed Expected Results"]))
                self.assertIn(case.test_id, case.tags)
                self.assertIsNotNone(importer.priority_number(case.priority))

    def test_exact_and_normalized_sheet_names(self):
        for sheet_name in ("Test_Cases", "Test Cases", "test_cases"):
            with self.subTest(sheet_name=sheet_name):
                cases = importer.read_excel(WORKBOOK, sheet_name, "Override", False)
                self.assertEqual(len(cases), 14)
                self.assertEqual({case.module for case in cases}, {"Override"})

    def test_missing_sheet_lists_available_names(self):
        with self.assertRaisesRegex(ValueError, "Available sheets:.*Test_Cases"):
            importer.read_excel(WORKBOOK, "Does not exist", None, False)

    def test_bdd_sheet_maps_steps_and_results(self):
        cases = importer.read_excel(WORKBOOK, "BDD_Test_Cases", None, False)
        self.assertTrue(cases)
        self.assertTrue(all(case.test_id and case.actions and case.expected for case in cases))

    def test_priorities(self):
        for value, expected in {"P0": 1, "P1": 2, "p2": 3, "P3": 4,
                                "high": 1, "medium": 2, "low": 3, "4": 4}.items():
            with self.subTest(value=value):
                self.assertEqual(importer.priority_number(value), expected)

    def test_duplicate_excel_ids_are_rejected(self):
        workbook = Workbook()
        workbook.active.title = "Test_Cases"
        workbook.active.append(["TC_ID", "Scenario", "Feature"])
        workbook.active.append(["TC-001", "First scenario", "Module"])
        workbook.active.append([" tc-001 ", "Renamed scenario", "Module"])
        with patch.object(importer, "load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "Duplicate TC_ID.*rows 2 and 3"):
                importer.read_excel(WORKBOOK, "Test_Cases", None, False)

    def test_tc_id_column_has_precedence_over_generic_id(self):
        headers = {"tc_id": 0, "id": 1}
        self.assertEqual(importer.find_column(headers, importer.ALIASES["test_id"]), 0)

    def test_id_tag_matching_is_exact(self):
        self.assertEqual(importer.test_case_tags("TC-0010; smoke", "TC-001"), "TC-0010; smoke; TC-001")
        self.assertEqual(importer.test_case_tags("tc-001; smoke", "TC-001"), "tc-001; smoke")


class AzureRequestTests(unittest.TestCase):
    def setUp(self):
        self.client = importer.AzureDevOpsClient("example", "example", "test-token")
        self.addCleanup(self.client.session.close)
        self.case = importer.TestCaseRow(
            excel_row=2, test_id="TC-001", title="Example", module="Module",
            preconditions="Signed in", actions=["Open page", "Submit"],
            expected=["Success"], priority="P1", tags="TC-001", description="Sample data",
        )

    def test_create_uses_post_and_json_patch_fields(self):
        with patch.object(self.client, "request", side_effect=[{"workItems": []}, {"id": 123}]) as request:
            with redirect_stdout(io.StringIO()):
                self.assertEqual(self.client.create_test_case(self.case), 123)
        self.assertEqual(request.call_args.args, ("POST", "/_apis/wit/workitems/$Test%20Case"))
        self.assertEqual(request.call_args.kwargs["content_type"], "application/json-patch+json")
        fields = {field["path"]: field["value"] for field in request.call_args.kwargs["json_body"]}
        self.assertEqual(fields["/fields/System.Title"], "[TC-001] Example")
        self.assertEqual(fields["/fields/Microsoft.VSTS.Common.Priority"], 2)
        steps = ElementTree.fromstring(fields["/fields/Microsoft.VSTS.TCM.Steps"])
        self.assertEqual(len(steps.findall("step")), 2)

    def test_existing_id_updates_even_when_title_changes(self):
        responses = [
            {"workItems": [{"id": 456}]},
            {"value": [{"id": 456, "fields": {"System.Title": "[TC-001] Old title", "System.Tags": "TC-001"}}]},
            {"id": 456},
        ]
        output = io.StringIO()
        with patch.object(self.client, "request", side_effect=responses) as request:
            with redirect_stdout(output):
                self.assertEqual(self.client.create_test_case(self.case), 456)
        self.assertEqual(request.call_args.args, ("PATCH", "/_apis/wit/workitems/456"))
        self.assertEqual(request.call_args.kwargs["content_type"], "application/json-patch+json")
        fields = {field["path"]: field["value"] for field in request.call_args.kwargs["json_body"]}
        self.assertEqual(fields["/fields/System.Title"], "[TC-001] Example")
        self.assertEqual(fields["/fields/Microsoft.VSTS.TCM.Steps"], importer.build_steps_xml(self.case.actions, self.case.expected))
        self.assertEqual(fields["/fields/Microsoft.VSTS.Common.Priority"], 2)
        self.assertIn("Signed in", fields["/fields/System.Description"])
        self.assertIn("Sample data", fields["/fields/System.Description"])
        self.assertEqual(fields["/fields/System.Tags"], "TC-001")
        self.assertIn("[UPDATED]", output.getvalue())

    def test_id_candidates_require_exact_tag_or_title_prefix(self):
        candidates = [
            {"id": 1, "fields": {"System.Title": "[TC-0010] Other", "System.Tags": "TC-0010"}},
            {"id": 2, "fields": {"System.Title": "Unrelated mention [TC-001]", "System.Tags": "prefix-TC-001"}},
            {"id": 3, "fields": {"System.Title": "Renamed by user", "System.Tags": "smoke; tc-001"}},
        ]
        with patch.object(self.client, "request", side_effect=[
            {"workItems": [{"id": item["id"]} for item in candidates]}, {"value": candidates}
        ]) as request:
            self.assertEqual(self.client.find_test_case_by_id("TC-001"), 3)
        query = request.call_args_list[0].kwargs["json_body"]["query"]
        self.assertIn("[System.TeamProject] = @project", query)
        self.assertIn("[System.Tags] CONTAINS 'TC-001'", query)

    def test_legacy_title_prefix_is_found_without_tag(self):
        with patch.object(self.client, "request", side_effect=[
            {"workItems": [{"id": 456}]},
            {"value": [{"id": 456, "fields": {"System.Title": "[TC-001] Legacy case"}}]},
        ]):
            self.assertEqual(self.client.find_test_case_by_id("TC-001"), 456)

    def test_ambiguous_ids_do_not_modify_any_cases(self):
        candidates = [{"id": case_id, "fields": {"System.Tags": "TC-001"}} for case_id in (1, 2)]
        with patch.object(self.client, "request", side_effect=[
            {"workItems": [{"id": 1}, {"id": 2}]}, {"value": candidates}
        ]) as request:
            with self.assertRaisesRegex(importer.AzureDevOpsError, "Multiple Test Cases have TC_ID"):
                self.client.create_test_case(self.case)
        self.assertEqual(request.call_count, 2)

    def test_id_query_escapes_quotes(self):
        with patch.object(self.client, "request", return_value={"workItems": []}) as request:
            self.assertIsNone(self.client.find_test_case_by_id("TC'O1"))
        self.assertIn("CONTAINS 'TC''O1'", request.call_args.kwargs["json_body"]["query"])

    def test_no_id_falls_back_to_title_for_update(self):
        self.case.test_id = ""
        with patch.object(self.client, "request", side_effect=[{"workItems": [{"id": 456}]}, {"id": 456}]) as request:
            with redirect_stdout(io.StringIO()):
                self.assertEqual(self.client.create_test_case(self.case), 456)
        self.assertEqual(request.call_args.args, ("PATCH", "/_apis/wit/workitems/456"))
        self.assertIn("[System.Title] = 'Example'", request.call_args_list[0].kwargs["json_body"]["query"])

    def test_repeated_import_creates_once_then_updates_same_id(self):
        stored = {}
        writes = []

        def request(method, path, **kwargs):
            if path == "/_apis/wit/wiql":
                return {"workItems": [{"id": 456}] if stored else []}
            if method == "GET":
                return {"value": [{"id": 456, "fields": stored}]}
            writes.append((method, path))
            stored.update({field["path"].removeprefix("/fields/"): field["value"]
                           for field in kwargs["json_body"]})
            return {"id": 456}

        with patch.object(self.client, "request", side_effect=request), redirect_stdout(io.StringIO()):
            self.assertEqual(self.client.create_test_case(self.case), 456)
            self.case.title = "Changed scenario"
            self.case.actions = ["New action"]
            self.assertEqual(self.client.create_test_case(self.case), 456)
        self.assertEqual(writes, [("POST", "/_apis/wit/workitems/$Test%20Case"), ("PATCH", "/_apis/wit/workitems/456")])
        self.assertEqual(stored["System.Title"], "[TC-001] Changed scenario")
        self.assertIn("New action", stored["Microsoft.VSTS.TCM.Steps"])

    def test_steps_do_not_mutate_input(self):
        importer.build_steps_xml(self.case.actions, self.case.expected)
        self.assertEqual(self.case.actions, ["Open page", "Submit"])
        self.assertEqual(self.case.expected, ["Success"])

    def test_request_reports_redirects_auth_errors_and_non_json(self):
        for status, body in ((302, ""), (401, "Unauthorized"), (403, "Forbidden"),
                             (200, "<html>Sign in</html>"), (200, "not JSON")):
            with self.subTest(status=status, body=body):
                response = importer.requests.Response()
                response.status_code = status
                response._content = body.encode()
                response.headers["Content-Type"] = "text/html"
                with patch.object(self.client.session, "request", return_value=response) as request:
                    with self.assertRaisesRegex(importer.AzureDevOpsError, "AZURE_DEVOPS_PAT"):
                        self.client.list_test_plans()
                self.assertFalse(request.call_args.kwargs["allow_redirects"])

    def test_request_accepts_json_and_empty_responses(self):
        for status, body, expected in ((200, b'{"value": []}', {"value": []}), (204, b"", None)):
            with self.subTest(status=status):
                response = importer.requests.Response()
                response.status_code = status
                response._content = body
                with patch.object(self.client.session, "request", return_value=response):
                    self.assertEqual(self.client.request("GET", "/_apis/testplan/plans"), expected)


class CommandLineTests(unittest.TestCase):
    def test_invalid_excel_stops_before_azure_changes(self):
        with patch.dict(os.environ, {"AZURE_DEVOPS_PAT": "test-token"}), patch("sys.argv", ["importer", str(WORKBOOK)]):
            with patch.object(importer, "read_excel", side_effect=ValueError("Duplicate TC_ID")):
                with patch.object(importer, "AzureDevOpsClient") as client, redirect_stderr(io.StringIO()):
                    self.assertEqual(importer.main(), 2)
                    client.assert_not_called()

    def test_dry_run_needs_no_token_or_network(self):
        output = io.StringIO()
        with patch.dict(os.environ, {}, clear=True), patch("sys.argv", ["importer", str(WORKBOOK), "--dry-run"]):
            with patch.object(importer, "AzureDevOpsClient") as client, redirect_stdout(output):
                self.assertEqual(importer.main(), 0)
                client.assert_not_called()
        self.assertIn("Read 14 test case(s)", output.getvalue())

    def test_live_import_requires_environment_token(self):
        error = io.StringIO()
        with patch.dict(os.environ, {}, clear=True), patch("sys.argv", ["importer", str(WORKBOOK)]):
            with patch.object(importer, "AzureDevOpsClient") as client, redirect_stderr(error):
                self.assertEqual(importer.main(), 2)
                client.assert_not_called()
        self.assertIn("Set AZURE_DEVOPS_PAT", error.getvalue())

    def test_plan_request_failure_stops_import_cleanly(self):
        error = io.StringIO()
        with patch.dict(os.environ, {"AZURE_DEVOPS_PAT": "test-token"}), patch("sys.argv", ["importer", str(WORKBOOK)]):
            with patch.object(importer, "AzureDevOpsClient") as client_type:
                client = client_type.return_value
                client.get_or_create_plan.side_effect = importer.AzureDevOpsError("Authentication required")
                with redirect_stdout(io.StringIO()), redirect_stderr(error):
                    self.assertEqual(importer.main(), 1)
                client.get_or_create_suite.assert_not_called()
                client.create_test_case.assert_not_called()
        self.assertEqual(error.getvalue().strip(), "ERROR: Authentication required")

    def test_import_flow_uses_environment_token_and_links_cases(self):
        with patch.dict(os.environ, {"AZURE_DEVOPS_PAT": "test-token"}), patch("sys.argv", ["importer", str(WORKBOOK)]):
            with patch.object(importer, "AzureDevOpsClient") as client_type, patch.object(importer.time, "sleep"):
                client = client_type.return_value
                client.get_or_create_plan.return_value = {"id": 100}
                client.get_or_create_suite.return_value = {"id": 200}
                client.create_test_case.side_effect = range(1, 15)
                with redirect_stdout(io.StringIO()):
                    self.assertEqual(importer.main(), 0)
                client_type.assert_called_once_with(importer.DEFAULT_ORG, importer.DEFAULT_PROJECT, "test-token")
                self.assertEqual(client.create_test_case.call_count, 14)
                client.add_test_cases_to_suite.assert_called_once_with(100, 200, list(range(1, 15)))


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        path = Path(__file__).resolve().parents[2] / ".github/workflows/upload-test-plan-to-azure.yml"
        self.workflow = yaml.load(path.read_text(encoding="utf-8"), Loader=yaml.BaseLoader)
        self.steps = {step["name"]: step for step in self.workflow["jobs"]["upload"]["steps"]}
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.workspace = Path(directory.name).resolve()
        self.excel = self.workspace / "uploads/cases with spaces.xlsx"
        self.excel.parent.mkdir()
        self.excel.touch()
        self.env = {
            "GITHUB_WORKSPACE": str(self.workspace),
            "RUNNER_TEMP": str(self.workspace),
            "GITHUB_STEP_SUMMARY": str(self.workspace / "summary.md"),
            "GITHUB_REPOSITORY": "example/tests",
            "GITHUB_ACTOR": "tester",
            "EXCEL_PATH": "uploads/cases with spaces.xlsx",
            "UPLOAD_ID": "",
            "PLAN_NAME": "Visits; $(echo not-a-shell-command)",
            "SHEET_NAME": "Test Cases",
            "AZURE_ORGANIZATION": "example",
            "AZURE_PROJECT": "example",
            "SUITE_NAME": "Visit Summaries",
        }

    def execute_step(self, name):
        mode, operation = {
            "Validate workbook": ("validate", workflow_runner.validate_workbook),
            "Upload test plan to Azure": ("upload", workflow_runner.upload_test_plan),
            "Delete temporary workbook upload": ("cleanup", workflow_runner.cleanup_workbook),
            "Write run summary": ("summary", workflow_runner.write_summary),
        }[name]
        self.assertEqual(self.steps[name]["run"], f"python tools/azure_test_importer/workflow_runner.py {mode}")
        operation()

    def test_manual_inputs_and_secret_scope(self):
        self.assertEqual(self.workflow["name"], "Upload Test Plan to Azure")
        self.assertEqual(set(self.workflow["on"]), {"workflow_dispatch"})
        inputs = self.workflow["on"]["workflow_dispatch"]["inputs"]
        self.assertTrue({"excel_path", "plan_name", "sheet_name"} <= inputs.keys())
        self.assertEqual(inputs["dry_run"]["default"], "true")
        self.assertNotIn("AZURE_DEVOPS_PAT", self.workflow["jobs"]["upload"]["env"])
        self.assertEqual(self.steps["Upload test plan to Azure"]["if"], "${{ !inputs.dry_run }}")
        self.assertEqual(self.steps["Upload test plan to Azure"]["env"]["AZURE_DEVOPS_PAT"],
                 "${{ secrets[format('AZURE_IMPORT_PAT_{0}', inputs.upload_id)] }}")
        self.assertEqual(self.workflow["permissions"], {"contents": "read"})
        self.assertEqual(self.workflow["jobs"]["upload"]["permissions"], {"contents": "write"})
        self.assertEqual(inputs["excel_path"]["required"], "false")
        self.assertEqual(inputs["upload_id"]["required"], "false")
        self.assertIn("always()", self.steps["Delete temporary workbook upload"]["if"])
        self.assertIn("GH_TOKEN", self.steps["Validate workbook"]["env"])
        self.assertNotIn("AZURE_DEVOPS_PAT", self.steps["Validate workbook"]["env"])
        self.assertEqual(self.workflow["concurrency"]["cancel-in-progress"], "false")
        for step in self.steps.values():
            self.assertNotIn("${{", step.get("run", ""))

    def test_validation_preserves_paths_and_inputs_as_arguments(self):
        with patch.dict(os.environ, self.env, clear=True), patch.object(importer, "main", return_value=0) as import_main:
            self.execute_step("Validate workbook")
        arguments = import_main.call_args.args[0]
        self.assertEqual(arguments[0], str(self.excel))
        self.assertEqual(arguments[arguments.index("--plan") + 1], self.env["PLAN_NAME"])
        self.assertEqual(arguments[arguments.index("--sheet") + 1], "Test Cases")
        self.assertEqual(arguments[-3:], ["--suite", "Visit Summaries", "--dry-run"])
        import_main.assert_called_once_with(arguments)
        saved = json.loads((self.workspace / "azure-import-arguments.json").read_text(encoding="utf-8"))
        self.assertEqual(saved, arguments[:-1])

    def test_invalid_paths_never_start_importer(self):
        for excel_path in ("../outside.xlsx", str(self.excel), "missing.xlsx", "cases.csv",
                   r"C:\Users\Tester\Downloads\cases.xlsx"):
            with self.subTest(excel_path=excel_path):
                with patch.dict(os.environ, {**self.env, "EXCEL_PATH": excel_path}, clear=True):
                    with patch.object(importer, "main") as import_main, self.assertRaises(SystemExit):
                        self.execute_step("Validate workbook")
                    import_main.assert_not_called()

    def test_upload_id_and_repository_path_are_mutually_exclusive(self):
        for upload_id, excel_path in (("", ""), ("a" * 32, "cases.xlsx")):
            with self.subTest(upload_id=upload_id, excel_path=excel_path):
                with patch.dict(os.environ, {**self.env, "UPLOAD_ID": upload_id, "EXCEL_PATH": excel_path}, clear=True):
                    with patch.object(importer, "main") as import_main, self.assertRaisesRegex(SystemExit, "either"):
                        self.execute_step("Validate workbook")
                    import_main.assert_not_called()

    def test_live_repository_path_requires_personal_uploader(self):
        with patch.dict(os.environ, {**self.env, "VALIDATION_ONLY": "false"}, clear=True):
            with patch.object(importer, "main") as import_main, self.assertRaisesRegex(SystemExit, "personal PAT"):
                self.execute_step("Validate workbook")
            import_main.assert_not_called()

    def test_uploaded_workbook_is_downloaded_and_validated_before_live_import(self):
        environment = {**self.env, "EXCEL_PATH": "", "UPLOAD_ID": "a" * 32}
        with patch.dict(os.environ, environment, clear=True):
            with patch.object(local_upload, "download_workbook", return_value=self.excel) as download:
                with patch.object(importer, "main", return_value=0) as import_main:
                    self.execute_step("Validate workbook")
        download.assert_called_once_with("example/tests", "a" * 32, "tester", str(self.workspace))
        self.assertEqual(import_main.call_args.args[0][0], str(self.excel))
        self.assertEqual(import_main.call_args.args[0][-1], "--dry-run")

    def test_failed_download_never_starts_importer(self):
        environment = {**self.env, "EXCEL_PATH": "", "UPLOAD_ID": "a" * 32}
        with patch.dict(os.environ, environment, clear=True):
            with patch.object(local_upload, "download_workbook", side_effect=ValueError("Checksum mismatch")):
                with patch.object(importer, "main") as import_main, self.assertRaisesRegex(ValueError, "Checksum"):
                    self.execute_step("Validate workbook")
                import_main.assert_not_called()

    def test_cleanup_is_scoped_to_initiating_user_and_uploaded_workbook(self):
        for upload_id in ("", "a" * 32):
            with self.subTest(upload_id=upload_id):
                with patch.dict(os.environ, {**self.env, "UPLOAD_ID": upload_id}, clear=True):
                    with patch.object(local_upload, "delete_upload") as delete, redirect_stdout(io.StringIO()):
                        self.execute_step("Delete temporary workbook upload")
                if upload_id:
                    delete.assert_called_once_with("example/tests", upload_id, "tester")
                else:
                    delete.assert_not_called()

    def test_validation_cli_runs_real_importer_without_pat(self):
        workspace = Path(__file__).resolve().parents[2]
        environment = {
            **os.environ, **self.env,
            "GITHUB_WORKSPACE": str(workspace),
            "EXCEL_PATH": "tools/azure_test_importer/example.xlsx",
            "SHEET_NAME": "Test_Cases",
        }
        environment.pop("AZURE_DEVOPS_PAT", None)
        result = subprocess.run(
            [sys.executable, str(Path(workflow_runner.__file__)), "validate"],
            env=environment, capture_output=True, text=True, check=True,
        )
        self.assertIn("Read 1 test case(s)", result.stdout)
        self.assertIn("No Azure DevOps changes were made.", result.stdout)
        saved = json.loads((self.workspace / "azure-import-arguments.json").read_text(encoding="utf-8"))
        self.assertNotIn("--dry-run", saved)

    def test_summary_preserves_mode_and_status(self):
        for validation_only, status, expected in (
            ("true", "success", "Excel validation only; no Azure changes"),
            ("false", "failure", "Azure import"),
        ):
            with self.subTest(validation_only=validation_only, status=status):
                environment = {**self.env, "VALIDATION_ONLY": validation_only, "JOB_STATUS": status}
                with patch.dict(os.environ, environment, clear=True):
                    self.execute_step("Write run summary")
                summary = Path(self.env["GITHUB_STEP_SUMMARY"]).read_text(encoding="utf-8")
                self.assertIn("Mode: " + expected, summary)
                self.assertIn("Status: " + status, summary)

    def test_failed_validation_stops_before_live_arguments_are_saved(self):
        with patch.dict(os.environ, self.env, clear=True):
            with patch.object(importer, "main", return_value=2):
                with self.assertRaises(SystemExit) as failure:
                    self.execute_step("Validate workbook")
        self.assertEqual(failure.exception.code, 2)
        self.assertFalse((self.workspace / "azure-import-arguments.json").exists())

    def test_missing_pat_prevents_upload(self):
        with patch.dict(os.environ, self.env, clear=True), patch.object(importer, "main") as import_main:
            with self.assertRaisesRegex(SystemExit, "Missing AZURE_DEVOPS_PAT"):
                self.execute_step("Upload test plan to Azure")
        import_main.assert_not_called()

    def test_upload_uses_environment_secret_not_command_argument(self):
        with patch.dict(os.environ, self.env, clear=True), patch.object(importer, "main", return_value=0):
            self.execute_step("Validate workbook")
        output = io.StringIO()
        with patch.dict(os.environ, {**self.env, "AZURE_DEVOPS_PAT": "test-token"}, clear=True):
            with patch.object(importer, "main", return_value=0) as import_main, redirect_stdout(output):
                self.execute_step("Upload test plan to Azure")
        self.assertNotIn("--dry-run", import_main.call_args.args[0])
        self.assertNotIn("test-token", " ".join(import_main.call_args.args[0]))
        import_main.assert_called_once_with(import_main.call_args.args[0])
        self.assertIn("::add-mask::test-token", output.getvalue())

    def test_failed_upload_propagates_exit_status(self):
        with patch.dict(os.environ, self.env, clear=True), patch.object(importer, "main", return_value=0):
            self.execute_step("Validate workbook")
        with patch.dict(os.environ, {**self.env, "AZURE_DEVOPS_PAT": "test-token"}, clear=True):
            with patch.object(importer, "main", return_value=1), redirect_stdout(io.StringIO()):
                with self.assertRaises(SystemExit) as failure:
                    self.execute_step("Upload test plan to Azure")
        self.assertEqual(failure.exception.code, 1)

    def test_multiline_pat_is_rejected_before_masking(self):
        output = io.StringIO()
        with patch.dict(os.environ, {**self.env, "AZURE_DEVOPS_PAT": "test\ntoken"}, clear=True):
            with patch.object(importer, "main") as import_main, redirect_stdout(output):
                with self.assertRaisesRegex(SystemExit, "line breaks"):
                    self.execute_step("Upload test plan to Azure")
        import_main.assert_not_called()
        self.assertEqual(output.getvalue(), "")


class LocalUploadTests(unittest.TestCase):
    def setUp(self):
        self.upload_id = "a" * 32
        self.release = {
            "id": 123, "draft": True, "tag_name": local_upload.upload_tag(self.upload_id),
            "author": {"login": "tester"},
            "body": json.dumps({"sha256": local_upload.validate_file(WORKBOOK)}),
            "assets": [{"name": "workbook.xlsx", "state": "uploaded", "size": WORKBOOK.stat().st_size}],
        }

    def test_picker_cancellation_makes_no_github_request(self):
        with patch.object(local_upload, "choose_file", return_value=None):
            with patch.object(local_upload, "github_json") as github, redirect_stdout(io.StringIO()):
                self.assertEqual(local_upload.main([]), 0)
        github.assert_not_called()

    def test_local_validation_defaults_to_dry_run(self):
        with patch.object(local_upload, "publish_workbook") as publish, redirect_stdout(io.StringIO()):
            self.assertEqual(local_upload.main(["--file", str(WORKBOOK), "--plan", "Visits", "--yes"]), 0)
        inputs = publish.call_args.args[3]
        self.assertEqual(inputs["dry_run"], "true")
        self.assertEqual(inputs["plan_name"], "AUTO_Visits")
        self.assertNotIn("AZURE_DEVOPS_PAT", inputs)

    def test_live_mode_is_explicit(self):
        with patch.object(local_upload, "publish_workbook") as publish, redirect_stdout(io.StringIO()):
            with patch.object(local_upload, "prompt_personal_pat", return_value="personal-test-token") as prompt:
                self.assertEqual(local_upload.main(["--file", str(WORKBOOK), "--live", "--yes"]), 0)
        self.assertEqual(publish.call_args.args[3]["dry_run"], "false")
        self.assertEqual(publish.call_args.kwargs["pat"], "personal-test-token")
        prompt.assert_called_once()

    def test_pat_cancellation_never_uploads(self):
        with patch.object(local_upload, "publish_workbook") as publish, redirect_stdout(io.StringIO()):
            with patch.object(local_upload, "prompt_personal_pat", return_value=None):
                self.assertEqual(local_upload.main(["--file", str(WORKBOOK), "--live", "--yes"]), 0)
        publish.assert_not_called()

    def test_dry_run_never_prompts_for_pat(self):
        with patch.object(local_upload, "publish_workbook"), redirect_stdout(io.StringIO()):
            with patch.object(local_upload, "prompt_personal_pat") as prompt:
                self.assertEqual(local_upload.main(["--file", str(WORKBOOK), "--yes"]), 0)
        prompt.assert_not_called()

    def test_live_pat_is_sent_on_stdin_and_deleted_after_workflow(self):
        for failure in (None, RuntimeError("Workflow failed")):
            with self.subTest(failure=failure):
                output = io.StringIO()
                with patch.object(local_upload, "github_json", return_value={"private": True, "default_branch": "main"}):
                    with patch.object(local_upload, "run_gh") as github, patch.object(local_upload, "dispatch_workbook"):
                        with patch.object(local_upload, "wait_for_workflow", side_effect=failure), redirect_stdout(output), redirect_stderr(output):
                            if failure:
                                with self.assertRaisesRegex(RuntimeError, "Workflow failed"):
                                    local_upload.publish_workbook(WORKBOOK, "example/tests", "main", {"dry_run": "false"}, pat="personal-test-token")
                            else:
                                local_upload.publish_workbook(WORKBOOK, "example/tests", "main", {"dry_run": "false"}, pat="personal-test-token")
                setter, deleter = github.call_args_list
                self.assertEqual(setter.args[0][:2], ["secret", "set"])
                self.assertEqual(setter.kwargs["stdin"], "personal-test-token")
                self.assertTrue(setter.kwargs["sensitive"])
                self.assertEqual(deleter.args[0][:2], ["secret", "delete"])
                self.assertEqual(setter.args[0][2], deleter.args[0][2])
                self.assertNotIn("personal-test-token", " ".join(setter.args[0]))
                self.assertNotIn(setter.args[0][2], output.getvalue())
                self.assertNotIn("personal-test-token", output.getvalue())

    def test_cleanup_warning_does_not_log_credentials_or_secret_name(self):
        output = io.StringIO()
        with patch.object(local_upload, "github_json", return_value={"private": True, "default_branch": "main"}):
            with patch.object(local_upload, "run_gh", side_effect=[None, RuntimeError("personal-test-token")]) as github:
                with patch.object(local_upload, "dispatch_workbook"), patch.object(local_upload, "wait_for_workflow"):
                    with redirect_stdout(output), redirect_stderr(output):
                        local_upload.publish_workbook(WORKBOOK, "example/tests", "main", {"dry_run": "false"}, pat="personal-test-token")
        self.assertIn("cleanup failed", output.getvalue())
        self.assertIn("revoke the PAT", output.getvalue())
        self.assertNotIn(github.call_args_list[0].args[0][2], output.getvalue())
        self.assertNotIn("personal-test-token", output.getvalue())

    def test_live_without_pat_stops_before_github_calls(self):
        with patch.object(local_upload, "github_json") as github:
            with self.assertRaisesRegex(ValueError, "personal PAT"):
                local_upload.publish_workbook(WORKBOOK, "example/tests", "main", {"dry_run": "false"})
        github.assert_not_called()

    def test_live_pat_requires_default_branch(self):
        with patch.object(local_upload, "github_json", return_value={"private": True, "default_branch": "main"}):
            with patch.object(local_upload, "run_gh") as github, self.assertRaisesRegex(ValueError, "default branch"):
                local_upload.publish_workbook(WORKBOOK, "example/tests", "untrusted", {"dry_run": "false"}, pat="test-token")
        github.assert_not_called()

    def test_sensitive_cli_failure_does_not_echo_pat(self):
        failure = subprocess.CalledProcessError(1, ["gh"], stderr="echoed-test-token")
        with patch.object(local_upload.subprocess, "run", side_effect=failure):
            with self.assertRaisesRegex(RuntimeError, "secret-management") as error:
                local_upload.run_gh(["secret", "set", "TEST"], stdin="echoed-test-token", sensitive=True)
        self.assertNotIn("echoed-test-token", str(error.exception))

    def test_invalid_workbook_never_uploads(self):
        with patch.object(local_upload, "publish_workbook") as publish, redirect_stderr(io.StringIO()):
            self.assertEqual(local_upload.main(["--file", "missing.xlsx", "--yes"]), 1)
        publish.assert_not_called()

    def test_invalid_sheet_never_uploads(self):
        with patch.object(local_upload, "publish_workbook") as publish, redirect_stderr(io.StringIO()):
            self.assertEqual(local_upload.main(["--file", str(WORKBOOK), "--sheet", "Missing", "--yes"]), 2)
        publish.assert_not_called()

    def test_declined_confirmation_never_uploads(self):
        with patch.object(local_upload, "confirm_upload", return_value=False):
            with patch.object(local_upload, "publish_workbook") as publish, redirect_stdout(io.StringIO()):
                self.assertEqual(local_upload.main(["--file", str(WORKBOOK)]), 0)
        publish.assert_not_called()

    def test_public_repository_is_rejected_before_release_creation(self):
        with patch.object(local_upload, "github_json", return_value={"private": False}) as github:
            with self.assertRaisesRegex(ValueError, "private"):
                local_upload.publish_workbook(WORKBOOK, "example/tests", "main", {})
        github.assert_called_once()

    def test_draft_upload_and_dispatch_have_no_file_content_or_pat_in_inputs(self):
        with patch.object(local_upload, "github_json", side_effect=[{"private": True}, {"id": 123}, None]) as github:
            with patch.object(local_upload, "run_gh") as upload, redirect_stdout(io.StringIO()):
                upload_id = local_upload.publish_workbook(WORKBOOK, "example/tests", "main", {"dry_run": "true"})
        draft = github.call_args_list[1].kwargs["body"]
        self.assertIs(draft["draft"], True)
        self.assertEqual(json.loads(draft["body"])["sha256"], local_upload.validate_file(WORKBOOK))
        self.assertEqual(upload.call_args.args[0][:3], ["release", "upload", local_upload.upload_tag(upload_id)])
        payload = github.call_args_list[2].kwargs["body"]
        self.assertEqual(payload, {"ref": "main", "inputs": {"dry_run": "true", "excel_path": "", "upload_id": upload_id}})

    def test_failed_asset_upload_deletes_only_created_draft(self):
        with patch.object(local_upload, "github_json", side_effect=[{"private": True}, {"id": 123}, None]) as github:
            with patch.object(local_upload, "run_gh", side_effect=RuntimeError("Upload failed")):
                with self.assertRaisesRegex(RuntimeError, "Upload failed"):
                    local_upload.publish_workbook(WORKBOOK, "example/tests", "main", {})
        self.assertEqual(github.call_args.args, ("repos/example/tests/releases/123",))
        self.assertEqual(github.call_args.kwargs["method"], "DELETE")

    def test_ambiguous_dispatch_failure_retains_upload_for_possible_run(self):
        with patch.object(local_upload, "github_json", side_effect=[
            {"private": True}, {"id": 123}, RuntimeError("Network timeout"),
        ]) as github:
            with patch.object(local_upload, "run_gh"):
                with self.assertRaisesRegex(RuntimeError, "Check Actions before retrying"):
                    local_upload.publish_workbook(WORKBOOK, "example/tests", "main", {})
        self.assertFalse(any(call.kwargs.get("method") == "DELETE" for call in github.call_args_list))

    def test_invalid_upload_ids_are_rejected_before_github_calls(self):
        for upload_id in ("../other", "invalid", "a" * 33):
            with self.subTest(upload_id=upload_id), patch.object(local_upload, "github_json") as github:
                with self.assertRaisesRegex(ValueError, "upload ID"):
                    local_upload.delete_upload("example/tests", upload_id, "tester")
                github.assert_not_called()

    def test_foreign_or_published_release_is_never_deleted(self):
        for changes in ({"draft": False}, {"author": {"login": "other-user"}}):
            with self.subTest(changes=changes):
                with patch.object(local_upload, "github_json", return_value=[{**self.release, **changes}]) as github:
                    with self.assertRaisesRegex(ValueError, "owned"):
                        local_upload.delete_upload("example/tests", self.upload_id, "tester")
                self.assertEqual(github.call_count, 1)

    def test_cleanup_deletes_only_verified_release(self):
        with patch.object(local_upload, "github_json", side_effect=[[self.release], None]) as github:
            local_upload.delete_upload("example/tests", self.upload_id, "tester")
        self.assertEqual(github.call_args.args, ("repos/example/tests/releases/123",))
        self.assertEqual(github.call_args.kwargs["method"], "DELETE")

    def test_download_verifies_workbook_checksum(self):
        def download(arguments, body=None):
            destination = Path(arguments[arguments.index("--dir") + 1], "workbook.xlsx")
            shutil.copyfile(WORKBOOK, destination)

        for checksum, valid in ((local_upload.validate_file(WORKBOOK), True), ("0" * 64, False)):
            with self.subTest(valid=valid), tempfile.TemporaryDirectory() as directory:
                release = {**self.release, "body": json.dumps({"sha256": checksum})}
                with patch.object(local_upload, "github_json", return_value=[release]):
                    with patch.object(local_upload, "run_gh", side_effect=download):
                        if valid:
                            path = local_upload.download_workbook("example/tests", self.upload_id, "tester", directory)
                            self.assertEqual(path.read_bytes(), WORKBOOK.read_bytes())
                        else:
                            with self.assertRaisesRegex(ValueError, "checksum"):
                                local_upload.download_workbook("example/tests", self.upload_id, "tester", directory)

    def test_invalid_assets_are_rejected_before_download(self):
        for assets in ([], [{"name": "other.xlsx"}], [
            {"name": "workbook.xlsx", "state": "uploaded", "size": local_upload.MAX_BYTES + 1},
        ]):
            with self.subTest(assets=assets):
                with patch.object(local_upload, "github_json", return_value=[{**self.release, "assets": assets}]):
                    with patch.object(local_upload, "run_gh") as download, self.assertRaisesRegex(ValueError, "exactly one"):
                        local_upload.download_workbook("example/tests", self.upload_id, "tester", tempfile.gettempdir())
                    download.assert_not_called()

    def test_github_subprocess_does_not_receive_azure_pat(self):
        with patch.dict(os.environ, {"AZURE_DEVOPS_PAT": "test-token"}):
            with patch.object(local_upload.subprocess, "run", return_value=subprocess.CompletedProcess([], 0, "{}")) as process:
                local_upload.github_json("user")
        self.assertNotIn("AZURE_DEVOPS_PAT", process.call_args.kwargs["env"])
        self.assertNotIn("test-token", " ".join(process.call_args.args[0]))
        self.assertNotIn("shell", process.call_args.kwargs)

    def test_draft_lookup_paginates_and_ignores_unrelated_releases(self):
        unrelated = [{"tag_name": f"v{index}"} for index in range(100)]
        with patch.object(local_upload, "github_json", side_effect=[unrelated, [self.release]]) as github:
            self.assertEqual(local_upload.owned_release("example/tests", self.upload_id, "tester"), self.release)
        self.assertIn("page=2", github.call_args.args[0])

    def test_unrelated_release_is_never_deleted(self):
        with patch.object(local_upload, "github_json", return_value=[{**self.release, "tag_name": "v1.0"}]) as github:
            with self.assertRaisesRegex(ValueError, "not found"):
                local_upload.delete_upload("example/tests", self.upload_id, "tester")
        self.assertEqual(github.call_count, 1)

    def test_monitor_selects_only_this_upload_run(self):
        runs = [{"id": 1, "display_title": "Another import"},
                {"id": 2, "display_title": "Azure import - " + self.upload_id}]
        with patch.object(local_upload, "github_json", return_value={"workflow_runs": runs}):
            with patch.object(local_upload, "run_gh") as watch, redirect_stdout(io.StringIO()):
                local_upload.wait_for_workflow("example/tests", "main", self.upload_id)
        self.assertEqual(watch.call_args.args[0], ["run", "watch", "2", "--repo", "example/tests", "--exit-status"])

    def test_monitor_timeout_stops_without_selecting_another_run(self):
        with patch.object(local_upload.time, "monotonic", side_effect=[0, 0, 181]):
            with patch.object(local_upload.time, "sleep"), patch.object(local_upload, "github_json", return_value={"workflow_runs": []}):
                with patch.object(local_upload, "run_gh") as watch, self.assertRaisesRegex(RuntimeError, "not found"):
                    local_upload.wait_for_workflow("example/tests", "main", self.upload_id)
                watch.assert_not_called()


if __name__ == "__main__":
    unittest.main()