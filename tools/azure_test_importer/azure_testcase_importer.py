#!/usr/bin/env python3
"""Import Excel test cases into Azure DevOps Test Plans.

Reads one test case per Excel row, creates/fetches an AUTO_ Test Plan, and
creates/fetches static Test Suites by module. Test Cases are matched by TC_ID
within the project using an exact tag or the legacy [TC_ID] title prefix.
Existing cases are updated; missing cases are created and linked to the suite.
Without TC_ID, matching falls back to the exact title. Ambiguous matches and
duplicate Excel IDs are rejected. Blank optional fields are not cleared on
update; tags are replaced with the workbook tags plus TC_ID.

Required packages:
    py -m pip install requests openpyxl

PAT is read from environment variable AZURE_DEVOPS_PAT. Never put the PAT in
this source file or in the Excel workbook.
"""
from __future__ import annotations

import argparse
import base64
import html
import json
import os
import re
import sys
import time
from contextlib import closing
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import quote

import requests
from openpyxl import load_workbook

API_VERSION = "7.1"
DEFAULT_ORG = "ShaharDigital"
DEFAULT_PROJECT = "MedicalCenter"

ALIASES = {
    "test_id": ["TC_ID", "Test ID", "TestID", "ID"],
    "title": ["Scenario", "Title", "Test Case", "Test Case Title", "Name"],
    "module": ["Module", "Feature", "Test Suite", "Suite"],
    "preconditions": ["Preconditions", "Precondition", "Pre-Steps"],
    "action": ["Detailed Automation Steps", "Action", "Steps", "Test Steps", "Step", "When"],
    "expected": ["Detailed Expected Results", "Expected Result", "Expected", "Expected Results", "Then / Expected Result", "Then"],
    "priority": ["Priority"],
    "tags": ["Tags"],
    "automation_candidate": ["Automation Candidate"],
    "description": ["Description", "Test Data"],
}


@dataclass
class TestCaseRow:
    excel_row: int
    test_id: str
    title: str
    module: str
    preconditions: str
    actions: list[str]
    expected: list[str]
    priority: str
    tags: str
    description: str


class AzureDevOpsError(RuntimeError):
    pass


class AzureDevOpsClient:
    def __init__(self, organization: str, project: str, pat: str, dry_run: bool = False):
        self.organization = organization
        self.project = project
        self.base = f"https://dev.azure.com/{quote(organization, safe='')}/{quote(project, safe='')}"
        self.dry_run = dry_run
        token = base64.b64encode(f":{pat}".encode()).decode()
        self.session = requests.Session()
        self.session.headers.update({"Authorization": f"Basic {token}", "Accept": "application/json"})

    def request(self, method: str, path: str, *, params=None, json_body=None,
                content_type: str = "application/json") -> Any:
        url = f"{self.base}{path}"
        if self.dry_run and method.upper() in {"POST", "PATCH", "PUT", "DELETE"}:
            print(f"[DRY-RUN] {method.upper()} {url}")
            if json_body is not None:
                print(json.dumps(json_body, ensure_ascii=False, indent=2)[:3000])
            return {"id": -1, "name": "DRY-RUN"}
        headers = {"Content-Type": content_type}
        response = self.session.request(method, url, params=params, json=json_body,
                                        headers=headers, timeout=60, allow_redirects=False)
        if 300 <= response.status_code < 400:
            raise AzureDevOpsError(
                f"Azure DevOps redirected the API request (HTTP {response.status_code}). "
                "Check AZURE_DEVOPS_PAT and the organization/project URL; authentication may be required."
            )
        if response.status_code in {401, 403}:
            raise AzureDevOpsError(
                f"Azure DevOps returned HTTP {response.status_code}. Check AZURE_DEVOPS_PAT, "
                "organization/project access, and Work Items and Test Management Read & write permissions."
            )
        if not response.ok:
            try:
                detail = response.json()
            except ValueError:
                detail = response.text
            raise AzureDevOpsError(
                f"Azure DevOps returned HTTP {response.status_code} for {method} {response.url}\n"
                f"{json.dumps(detail, ensure_ascii=False, indent=2) if isinstance(detail, dict) else detail}"
            )
        if not response.content:
            return None
        try:
            return response.json()
        except ValueError as exc:
            raise AzureDevOpsError(
                f"Azure DevOps returned a non-JSON response (HTTP {response.status_code}, "
                f"Content-Type: {response.headers.get('Content-Type', 'unknown')}). "
                "Check AZURE_DEVOPS_PAT, the organization/project URL, and any proxy or sign-in page."
            ) from exc

    def list_test_plans(self) -> list[dict]:
        data = self.request("GET", "/_apis/testplan/plans", params={"api-version": API_VERSION})
        return data.get("value", [])

    def get_or_create_plan(self, name: str, area_path: str | None = None,
                           iteration: str | None = None) -> dict:
        name = normalize_plan_name(name)
        for plan in self.list_test_plans():
            if str(plan.get("name", "")).strip().casefold() == name.strip().casefold():
                print(f"[FOUND] Test Plan: {name} (ID {plan['id']})")
                return plan
        body: dict[str, Any] = {"name": name}
        if area_path:
            body["areaPath"] = area_path
        if iteration:
            body["iteration"] = iteration
        plan = self.request("POST", "/_apis/testplan/plans",
                            params={"api-version": API_VERSION}, json_body=body)
        print(f"[CREATED] Test Plan: {name} (ID {plan['id']})")
        return plan

    def list_suites(self, plan_id: int) -> list[dict]:
        data = self.request("GET", f"/_apis/testplan/Plans/{plan_id}/suites",
                            params={"asTreeView": "false", "api-version": API_VERSION})
        return data.get("value", [])

    def get_or_create_suite(self, plan: dict, suite_name: str) -> dict:
        plan_id = int(plan["id"])
        for suite in self.list_suites(plan_id):
            if str(suite.get("name", "")).strip().casefold() == suite_name.strip().casefold():
                print(f"[FOUND] Suite: {suite_name} (ID {suite['id']})")
                return suite
        root = plan.get("rootSuite") or {}
        root_id = root.get("id")
        if not root_id:
            refreshed = self.request("GET", f"/_apis/testplan/plans/{plan_id}",
                                     params={"api-version": API_VERSION})
            root_id = (refreshed.get("rootSuite") or {}).get("id")
        if not root_id:
            raise AzureDevOpsError(f"Could not determine root suite for Test Plan {plan_id}")
        body = {
            "suiteType": "staticTestSuite",
            "name": suite_name,
            "parentSuite": {"id": int(root_id)},
            "inheritDefaultConfigurations": True,
        }
        suite = self.request("POST", f"/_apis/testplan/Plans/{plan_id}/suites",
                             params={"api-version": API_VERSION}, json_body=body)
        print(f"[CREATED] Suite: {suite_name} (ID {suite['id']})")
        return suite

    def find_test_case_by_title(self, title: str) -> int | None:
        escaped = title.replace("'", "''")
        query = {
            "query": "SELECT [System.Id] FROM WorkItems "
                     "WHERE [System.TeamProject] = @project "
                     "AND [System.WorkItemType] = 'Test Case' "
                     f"AND [System.Title] = '{escaped}' ORDER BY [System.Id] DESC"
        }
        data = self.request("POST", "/_apis/wit/wiql",
                            params={"api-version": API_VERSION}, json_body=query)
        items = data.get("workItems", [])
        if len(items) > 1:
            raise AzureDevOpsError(f"Multiple Test Cases match title {title!r}; use a unique TC_ID.")
        return int(items[0]["id"]) if items else None

    def find_test_case_by_id(self, test_id: str) -> int | None:
        test_id = text(test_id)
        if not test_id:
            raise ValueError("TC_ID must not be empty.")
        escaped = test_id.replace("'", "''")
        query = {
            "query": "SELECT [System.Id] FROM WorkItems "
                     "WHERE [System.TeamProject] = @project "
                     "AND [System.WorkItemType] = 'Test Case' "
                     f"AND ([System.Tags] CONTAINS '{escaped}' "
                     f"OR [System.Title] CONTAINS '[{escaped}]')"
        }
        data = self.request("POST", "/_apis/wit/wiql",
                            params={"api-version": API_VERSION}, json_body=query)
        ids = [int(item["id"]) for item in data.get("workItems", [])]
        matches: set[int] = set()
        marker = f"[{test_id}]".casefold()
        for start in range(0, len(ids), 200):
            details = self.request("GET", "/_apis/wit/workitems", params={
                "api-version": API_VERSION,
                "ids": ",".join(map(str, ids[start:start + 200])),
                "fields": "System.Title,System.Tags",
            })
            for item in details["value"]:
                fields = item.get("fields", {})
                tags = {tag.strip().casefold() for tag in text(fields.get("System.Tags")).split(";")}
                title = text(fields.get("System.Title")).casefold()
                if test_id.casefold() in tags or title == marker or title.startswith(marker + " "):
                    matches.add(int(item["id"]))
        if len(matches) > 1:
            raise AzureDevOpsError(
                f"Multiple Test Cases have TC_ID {test_id!r}: {sorted(matches)}. "
                "Resolve the duplicates before importing."
            )
        return next(iter(matches), None)

    def create_test_case(self, tc: TestCaseRow) -> int:
        """Create or update a Test Case while retaining the existing public API."""
        test_id = text(tc.test_id)
        title = f"[{test_id}] {tc.title}" if test_id else tc.title
        existing = self.find_test_case_by_id(test_id) if test_id else self.find_test_case_by_title(title)

        fields: list[dict[str, Any]] = [
            {"op": "add", "path": "/fields/System.Title", "value": title},
            {"op": "add", "path": "/fields/Microsoft.VSTS.TCM.Steps", "value": build_steps_xml(tc.actions, tc.expected)},
        ]
        if tc.preconditions:
            fields.append({"op": "add", "path": "/fields/Microsoft.VSTS.TCM.LocalDataSource",
                           "value": f"<NewDataSet><Preconditions>{html.escape(tc.preconditions)}</Preconditions></NewDataSet>"})
        description_parts = []
        if tc.preconditions:
            description_parts.append(f"<b>Preconditions:</b><br>{html.escape(tc.preconditions)}")
        if tc.description:
            description_parts.append(f"<b>Test Data / Description:</b><br>{html.escape(tc.description)}")
        if description_parts:
            fields.append({"op": "add", "path": "/fields/System.Description",
                           "value": "<br><br>".join(description_parts)})
        tags = test_case_tags(tc.tags, test_id)
        if tags:
            fields.append({"op": "add", "path": "/fields/System.Tags", "value": tags})
        priority = priority_number(tc.priority)
        if priority:
            fields.append({"op": "add", "path": "/fields/Microsoft.VSTS.Common.Priority", "value": priority})

        result = self.request(
            "PATCH" if existing else "POST",
            f"/_apis/wit/workitems/{existing}" if existing else "/_apis/wit/workitems/$Test%20Case",
            params={"api-version": API_VERSION}, json_body=fields,
            content_type="application/json-patch+json",
        )
        operation = "UPDATED" if existing else "CREATED"
        print(f"[{operation}] Test Case: {title} (ID {result['id']})")
        return int(result["id"])

    def add_test_cases_to_suite(self, plan_id: int, suite_id: int, ids: Iterable[int]) -> None:
        ids = list(dict.fromkeys(case_id for case_id in ids if case_id > 0))
        if not ids:
            return
        id_path = ",".join(map(str, ids))
        self.request("POST", f"/_apis/test/Plans/{plan_id}/suites/{suite_id}/testcases/{id_path}",
                     params={"api-version": API_VERSION}, json_body=None)
        print(f"[LINKED] {len(ids)} Test Case(s) to Suite ID {suite_id}")


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def test_case_tags(value: str, test_id: str) -> str:
    tags = [tag.strip() for tag in value.split(";") if tag.strip()]
    if test_id and test_id.casefold() not in {tag.casefold() for tag in tags}:
        tags.append(test_id)
    return "; ".join(tags)


def normalize_plan_name(value: str) -> str:
    name = text(value)
    while name[:5].casefold() == "auto_":
        name = name[5:].strip()
    if not name:
        raise ValueError("Test Plan name must not be empty after AUTO_.")
    return f"AUTO_{name}"


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", text(value)).casefold()


def find_column(headers: dict[str, int], aliases: list[str]) -> int | None:
    for alias in aliases:
        if normalize_header(alias) in headers:
            return headers[normalize_header(alias)]
    return None


def split_multiline(value: str) -> list[str]:
    value = text(value)
    if not value:
        return []
    lines = [re.sub(r"^\s*(?:\d+[.)]|[-*•])\s*", "", x).strip()
             for x in re.split(r"\r?\n", value)]
    return [x for x in lines if x]


def align_steps(actions: list[str], expected: list[str]) -> tuple[list[str], list[str]]:
    if not actions:
        actions = ["Execute the test scenario"]
    if not expected:
        expected = ["The expected behavior is observed"]
    size = max(len(actions), len(expected))
    if len(actions) == 1 and size > 1:
        actions *= size
    if len(expected) == 1 and size > 1:
        expected *= size
    actions.extend([""] * (size - len(actions)))
    expected.extend([""] * (size - len(expected)))
    return actions, expected


def build_steps_xml(actions: list[str], expected: list[str]) -> str:
    actions, expected = align_steps(actions[:], expected[:])
    parts = [f'<steps id="0" last="{len(actions)}">']
    for idx, (action, exp) in enumerate(zip(actions, expected), start=1):
        a = html.escape(action)
        e = html.escape(exp)
        parts.append(
            f'<step id="{idx}" type="ActionStep">'
            f'<parameterizedString isformatted="true">&lt;DIV&gt;{a}&lt;/DIV&gt;</parameterizedString>'
            f'<parameterizedString isformatted="true">&lt;DIV&gt;{e}&lt;/DIV&gt;</parameterizedString>'
            f'<description/></step>'
        )
    parts.append("</steps>")
    return "".join(parts)


def priority_number(value: str) -> int | None:
    mapping = {"critical": 1, "highest": 1, "high": 1, "medium": 2, "normal": 2, "low": 3}
    mapping.update({"p0": 1, "p1": 2, "p2": 3, "p3": 4})
    if text(value).isdigit():
        return max(1, min(4, int(text(value))))
    return mapping.get(text(value).casefold())


def read_excel(path: Path, sheet_name: str, default_suite: str | None,
               only_automation_candidates: bool) -> list[TestCaseRow]:
    with closing(load_workbook(path, read_only=True, data_only=True)) as wb:
        if sheet_name not in wb.sheetnames:
            matches = [name for name in wb.sheetnames
                       if normalize_header(name.replace("_", " ")) == normalize_header(sheet_name.replace("_", " "))]
            if len(matches) != 1:
                raise ValueError(f"Sheet '{sheet_name}' not found or ambiguous. Available sheets: {', '.join(wb.sheetnames)}")
            sheet_name = matches[0]
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        try:
            first = next(rows)
        except StopIteration:
            return []
        headers = {normalize_header(v): idx for idx, v in enumerate(first) if text(v)}
        cols = {key: find_column(headers, aliases) for key, aliases in ALIASES.items()}
        if cols["title"] is None:
            raise ValueError(f"Missing title column. Supported names: {ALIASES['title']}")
        if cols["module"] is None and not default_suite:
            raise ValueError(f"Missing module/suite column. Supported names: {ALIASES['module']}; or pass --suite")

        def cell(row, key):
            idx = cols[key]
            return text(row[idx]) if idx is not None and idx < len(row) else ""

        output: list[TestCaseRow] = []
        seen_ids: dict[str, int] = {}
        for row_num, row in enumerate(rows, start=2):
            title = cell(row, "title")
            if not title:
                continue
            if only_automation_candidates and cell(row, "automation_candidate").casefold() not in {"yes", "true", "1", "y"}:
                continue
            module = default_suite or cell(row, "module") or "Imported Test Cases"
            tags = cell(row, "tags")
            test_id = cell(row, "test_id")
            if test_id:
                if test_id.casefold() in seen_ids:
                    raise ValueError(
                        f"Duplicate TC_ID {test_id!r} in Excel rows "
                        f"{seen_ids[test_id.casefold()]} and {row_num}."
                    )
                seen_ids[test_id.casefold()] = row_num
            tags = test_case_tags(tags, test_id)
            output.append(TestCaseRow(
                excel_row=row_num,
                test_id=test_id,
                title=title,
                module=module,
                preconditions=cell(row, "preconditions"),
                actions=split_multiline(cell(row, "action")),
                expected=split_multiline(cell(row, "expected")),
                priority=cell(row, "priority"),
                tags=tags,
                description=cell(row, "description"),
            ))
        return output


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Create or update Azure DevOps test cases from Excel by TC_ID")
    p.add_argument("excel", type=Path, help="Path to the .xlsx file")
    p.add_argument("--organization", default=DEFAULT_ORG)
    p.add_argument("--project", default=DEFAULT_PROJECT)
    p.add_argument("--plan", type=normalize_plan_name, default="Automation Test Cases",
                   help="Test Plan name (AUTO_ prefix is added automatically)")
    p.add_argument("--sheet", default="Test Cases", help="Excel sheet name")
    p.add_argument("--suite", help="Force all rows into one suite; otherwise Module/Feature is used")
    p.add_argument("--area-path", help="Optional areaPath when creating a new Test Plan")
    p.add_argument("--iteration", help="Optional iteration when creating a new Test Plan")
    p.add_argument("--only-automation-candidates", action="store_true")
    p.add_argument("--dry-run", action="store_true", help="Read and validate Excel without changing Azure DevOps")
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if not args.excel.exists():
        print(f"ERROR: Excel file not found: {args.excel}", file=sys.stderr)
        return 2
    pat = os.getenv("AZURE_DEVOPS_PAT", "").strip()
    if not pat and not args.dry_run:
        print("ERROR: Set AZURE_DEVOPS_PAT before running the script.", file=sys.stderr)
        return 2

    try:
        cases = read_excel(args.excel, args.sheet, args.suite, args.only_automation_candidates)
    except ValueError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    if not cases:
        print("No test cases were found after applying the selected filters.")
        return 0
    print(f"Read {len(cases)} test case(s) from {args.excel.name}")
    print(f"Test Plan: {args.plan}")
    suites: dict[str, list[TestCaseRow]] = {}
    for tc in cases:
        suites.setdefault(tc.module, []).append(tc)
    print("Suites:", ", ".join(f"{name} ({len(items)})" for name, items in suites.items()))
    if args.dry_run:
        print("Dry-run Excel validation completed. No Azure DevOps changes were made.")
        return 0

    client = AzureDevOpsClient(args.organization, args.project, pat)
    try:
        plan = client.get_or_create_plan(args.plan, args.area_path, args.iteration)
    except (AzureDevOpsError, requests.RequestException) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    failures: list[str] = []
    for suite_name, suite_cases in suites.items():
        try:
            suite = client.get_or_create_suite(plan, suite_name)
            ids = []
            for tc in suite_cases:
                try:
                    ids.append(client.create_test_case(tc))
                    time.sleep(0.05)
                except Exception as exc:
                    failures.append(f"Excel row {tc.excel_row} ({tc.test_id or tc.title}): {exc}")
            # Azure permits comma-separated IDs; chunk to keep URLs manageable.
            for start in range(0, len(ids), 100):
                client.add_test_cases_to_suite(int(plan["id"]), int(suite["id"]), ids[start:start + 100])
        except Exception as exc:
            failures.append(f"Suite '{suite_name}': {exc}")

    print(f"\nCompleted. Processed {len(cases)} Excel row(s); failures: {len(failures)}")
    if failures:
        print("\nFailures:", file=sys.stderr)
        for failure in failures:
            print(f"- {failure}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
