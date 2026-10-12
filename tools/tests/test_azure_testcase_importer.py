import io
import os
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook

from tools import azure_testcase_importer as importer

WORKBOOK = Path("Test_Cases.xlsx")


def setUpModule():
    global WORKBOOK
    temporary = TemporaryDirectory()
    unittest.addModuleCleanup(temporary.cleanup)
    WORKBOOK = Path(temporary.name) / "Test_Cases.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Test_Cases"
    worksheet.append([
        "TC_ID", "Scenario", "Feature", "Detailed Automation Steps",
        "Detailed Expected Results", "Priority", "Preconditions", "Test Data", "Tags",
    ])
    for index in range(1, 15):
        worksheet.append([
            f"TC-{index:03d}", f"Synthetic scenario {index}", "Synthetic module",
            "1. Open page\n2. Submit", "1. Page opens\n2. Success", "P1",
            "Signed in", "Synthetic data only", "smoke",
        ])
    bdd = workbook.create_sheet("BDD_Test_Cases")
    bdd.append(["TC_ID", "Scenario", "Feature", "When", "Then / Expected Result"])
    bdd.append(["BDD-001", "Synthetic BDD scenario", "Synthetic module", "Open page", "Page opens"])
    workbook.save(WORKBOOK)
    workbook.close()


class PlanNameTests(unittest.TestCase):
    def test_prefix_is_added_once(self):
        for name in ("Visits", "AUTO_Visits", "auto_Visits", " Auto_AUTO_Visits "):
            with self.subTest(name=name):
                self.assertEqual(importer.normalize_plan_name(name), "AUTO_Visits")

    def test_empty_name_is_rejected(self):
        for name in ("", "  ", "AUTO_", "AUTO_AUTO_"):
            with self.subTest(name=name), self.assertRaises(ValueError):
                importer.normalize_plan_name(name)

    def test_existing_mixed_case_prefix_is_reused(self):
        client = importer.AzureDevOpsClient("example", "example", "test-token")
        self.addCleanup(client.session.close)
        existing = {"id": 100, "name": "Auto_Visits"}
        with patch.object(client, "request", return_value={"value": [existing]}) as request:
            with redirect_stdout(io.StringIO()):
                self.assertEqual(client.get_or_create_plan("Visits"), existing)
        request.assert_called_once()

    def test_new_plan_has_uppercase_prefix(self):
        client = importer.AzureDevOpsClient("example", "example", "test-token")
        self.addCleanup(client.session.close)
        with patch.object(client, "request", side_effect=[{"value": []}, {"id": 100}]) as request:
            with redirect_stdout(io.StringIO()):
                client.get_or_create_plan("Visits")
        self.assertEqual(request.call_args.kwargs["json_body"]["name"], "AUTO_Visits")


class ExcelImportTests(unittest.TestCase):
    def test_workbook_preserves_detailed_columns(self):
        cases = importer.read_excel(WORKBOOK, "Test Cases", None, False)
        workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
        try:
            rows = workbook["Test_Cases"].iter_rows(values_only=True)
            headers = next(rows)
            records = [dict(zip(headers, row)) for row in rows]
        finally:
            workbook.close()
        records = [record for record in records if record["Scenario"]]
        self.assertEqual(len(cases), 14)
        self.assertEqual(len(cases), len(records))
        for case, record in zip(cases, records):
            with self.subTest(test_id=case.test_id):
                self.assertEqual(case.test_id, record["TC_ID"])
                self.assertEqual(case.title, record["Scenario"])
                self.assertEqual(case.module, record["Feature"])
                self.assertEqual(case.actions, importer.split_multiline(record["Detailed Automation Steps"]))
                self.assertEqual(case.expected, importer.split_multiline(record["Detailed Expected Results"]))
                self.assertIn(case.test_id, case.tags)
                self.assertIsNotNone(importer.priority_number(case.priority))

    def test_exact_and_normalized_sheet_names(self):
        for sheet_name in ("Test_Cases", "Test Cases", "test_cases"):
            with self.subTest(sheet_name=sheet_name):
                cases = importer.read_excel(WORKBOOK, sheet_name, "Override", False)
                self.assertEqual(len(cases), 14)
                self.assertEqual({case.module for case in cases}, {"Override"})

    def test_missing_sheet_lists_available_names(self):
        with self.assertRaisesRegex(ValueError, "Available sheets:.*Test_Cases"):
            importer.read_excel(WORKBOOK, "Does not exist", None, False)

    def test_bdd_sheet_maps_steps_and_results(self):
        cases = importer.read_excel(WORKBOOK, "BDD_Test_Cases", None, False)
        self.assertTrue(cases)
        self.assertTrue(all(case.test_id and case.actions and case.expected for case in cases))

    def test_priorities(self):
        for value, expected in {"P0": 1, "P1": 2, "p2": 3, "P3": 4,
                                "high": 1, "medium": 2, "low": 3, "4": 4}.items():
            with self.subTest(value=value):
                self.assertEqual(importer.priority_number(value), expected)

    def test_duplicate_excel_ids_are_rejected(self):
        workbook = Workbook()
        workbook.active.title = "Test_Cases"
        workbook.active.append(["TC_ID", "Scenario", "Feature"])
        workbook.active.append(["TC-001", "First scenario", "Module"])
        workbook.active.append([" tc-001 ", "Renamed scenario", "Module"])
        with patch.object(importer, "load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "Duplicate TC_ID.*rows 2 and 3"):
                importer.read_excel(WORKBOOK, "Test_Cases", None, False)

    def test_tc_id_column_has_precedence_over_generic_id(self):
        headers = {"tc_id": 0, "id": 1}
        self.assertEqual(importer.find_column(headers, importer.ALIASES["test_id"]), 0)

    def test_id_tag_matching_is_exact(self):
        self.assertEqual(importer.test_case_tags("TC-0010; smoke", "TC-001"), "TC-0010; smoke; TC-001")
        self.assertEqual(importer.test_case_tags("tc-001; smoke", "TC-001"), "tc-001; smoke")

    def test_workbook_closes_on_success_and_invalid_headers(self):
        for headers in (["TC_ID", "Scenario", "Feature"], ["Invalid"]):
            with self.subTest(headers=headers):
                workbook = Workbook()
                workbook.active.title = "Test_Cases"
                workbook.active.append(headers)
                with patch.object(importer, "load_workbook", return_value=workbook):
                    with patch.object(workbook, "close") as close:
                        if headers == ["Invalid"]:
                            with self.assertRaises(ValueError):
                                importer.read_excel(WORKBOOK, "Test_Cases", None, False)
                        else:
                            self.assertEqual(importer.read_excel(WORKBOOK, "Test_Cases", None, False), [])
                        close.assert_called_once()


class AzureRequestTests(unittest.TestCase):
    def setUp(self):
        self.client = importer.AzureDevOpsClient("example", "example", "test-token")
        self.addCleanup(self.client.session.close)
        self.case = importer.TestCaseRow(
            excel_row=2, test_id="TC-001", title="Example", module="Module",
            preconditions="Signed in", actions=["Open page", "Submit"],
            expected=["Success"], priority="P1", tags="TC-001", description="Sample data",
        )

    def test_create_uses_post_and_json_patch_fields(self):
        with patch.object(self.client, "request", side_effect=[{"workItems": []}, {"id": 123}]) as request:
            with redirect_stdout(io.StringIO()):
                self.assertEqual(self.client.create_test_case(self.case), 123)
        self.assertEqual(request.call_args.args, ("POST", "/_apis/wit/workitems/$Test%20Case"))
        self.assertEqual(request.call_args.kwargs["content_type"], "application/json-patch+json")
        fields = {field["path"]: field["value"] for field in request.call_args.kwargs["json_body"]}
        self.assertEqual(fields["/fields/System.Title"], "[TC-001] Example")
        self.assertEqual(fields["/fields/Microsoft.VSTS.Common.Priority"], 2)
        steps = ElementTree.fromstring(fields["/fields/Microsoft.VSTS.TCM.Steps"])
        self.assertEqual(len(steps.findall("step")), 2)

    def test_existing_id_updates_even_when_title_changes(self):
        responses = [
            {"workItems": [{"id": 456}]},
            {"value": [{"id": 456, "fields": {"System.Title": "[TC-001] Old title", "System.Tags": "TC-001"}}]},
            {"id": 456},
        ]
        output = io.StringIO()
        with patch.object(self.client, "request", side_effect=responses) as request:
            with redirect_stdout(output):
                self.assertEqual(self.client.create_test_case(self.case), 456)
        self.assertEqual(request.call_args.args, ("PATCH", "/_apis/wit/workitems/456"))
        self.assertEqual(request.call_args.kwargs["content_type"], "application/json-patch+json")
        fields = {field["path"]: field["value"] for field in request.call_args.kwargs["json_body"]}
        self.assertEqual(fields["/fields/System.Title"], "[TC-001] Example")
        self.assertEqual(fields["/fields/Microsoft.VSTS.TCM.Steps"], importer.build_steps_xml(self.case.actions, self.case.expected))
        self.assertEqual(fields["/fields/Microsoft.VSTS.Common.Priority"], 2)
        self.assertIn("Signed in", fields["/fields/System.Description"])
        self.assertIn("Sample data", fields["/fields/System.Description"])
        self.assertEqual(fields["/fields/System.Tags"], "TC-001")
        self.assertIn("[UPDATED]", output.getvalue())

    def test_id_candidates_require_exact_tag_or_title_prefix(self):
        candidates = [
            {"id": 1, "fields": {"System.Title": "[TC-0010] Other", "System.Tags": "TC-0010"}},
            {"id": 2, "fields": {"System.Title": "Unrelated mention [TC-001]", "System.Tags": "prefix-TC-001"}},
            {"id": 3, "fields": {"System.Title": "Renamed by user", "System.Tags": "smoke; tc-001"}},
        ]
        with patch.object(self.client, "request", side_effect=[
            {"workItems": [{"id": item["id"]} for item in candidates]}, {"value": candidates}
        ]) as request:
            self.assertEqual(self.client.find_test_case_by_id("TC-001"), 3)
        query = request.call_args_list[0].kwargs["json_body"]["query"]
        self.assertIn("[System.TeamProject] = @project", query)
        self.assertIn("[System.Tags] CONTAINS 'TC-001'", query)

    def test_legacy_title_prefix_is_found_without_tag(self):
        with patch.object(self.client, "request", side_effect=[
            {"workItems": [{"id": 456}]},
            {"value": [{"id": 456, "fields": {"System.Title": "[TC-001] Legacy case"}}]},
        ]):
            self.assertEqual(self.client.find_test_case_by_id("TC-001"), 456)

    def test_ambiguous_ids_do_not_modify_any_cases(self):
        candidates = [{"id": case_id, "fields": {"System.Tags": "TC-001"}} for case_id in (1, 2)]
        with patch.object(self.client, "request", side_effect=[
            {"workItems": [{"id": 1}, {"id": 2}]}, {"value": candidates}
        ]) as request:
            with self.assertRaisesRegex(importer.AzureDevOpsError, "Multiple Test Cases have TC_ID"):
                self.client.create_test_case(self.case)
        self.assertEqual(request.call_count, 2)

    def test_id_query_escapes_quotes(self):
        with patch.object(self.client, "request", return_value={"workItems": []}) as request:
            self.assertIsNone(self.client.find_test_case_by_id("TC'O1"))
        self.assertIn("CONTAINS 'TC''O1'", request.call_args.kwargs["json_body"]["query"])

    def test_no_id_falls_back_to_title_for_update(self):
        self.case.test_id = ""
        with patch.object(self.client, "request", side_effect=[{"workItems": [{"id": 456}]}, {"id": 456}]) as request:
            with redirect_stdout(io.StringIO()):
                self.assertEqual(self.client.create_test_case(self.case), 456)
        self.assertEqual(request.call_args.args, ("PATCH", "/_apis/wit/workitems/456"))
        self.assertIn("[System.Title] = 'Example'", request.call_args_list[0].kwargs["json_body"]["query"])

    def test_repeated_import_creates_once_then_updates_same_id(self):
        stored = {}
        writes = []

        def request(method, path, **kwargs):
            if path == "/_apis/wit/wiql":
                return {"workItems": [{"id": 456}]} if stored else {"workItems": []}
            if method == "GET":
                return {"value": [{"id": 456, "fields": stored}]}
            writes.append((method, path))
            stored.update({field["path"].removeprefix("/fields/"): field["value"]
                           for field in kwargs["json_body"]})
            return {"id": 456}

        with patch.object(self.client, "request", side_effect=request), redirect_stdout(io.StringIO()):
            self.assertEqual(self.client.create_test_case(self.case), 456)
            self.case.title = "Changed scenario"
            self.case.actions = ["New action"]
            self.assertEqual(self.client.create_test_case(self.case), 456)
        self.assertEqual(writes, [("POST", "/_apis/wit/workitems/$Test%20Case"), ("PATCH", "/_apis/wit/workitems/456")])
        self.assertEqual(stored["System.Title"], "[TC-001] Changed scenario")
        self.assertIn("New action", stored["Microsoft.VSTS.TCM.Steps"])

    def test_steps_do_not_mutate_input(self):
        importer.build_steps_xml(self.case.actions, self.case.expected)
        self.assertEqual(self.case.actions, ["Open page", "Submit"])
        self.assertEqual(self.case.expected, ["Success"])

    def test_request_reports_redirects_auth_errors_and_non_json(self):
        for status, body in ((302, ""), (401, "Unauthorized"), (403, "Forbidden"),
                             (200, "<html>Sign in</html>"), (200, "not JSON")):
            with self.subTest(status=status, body=body):
                response = importer.requests.Response()
                response.status_code = status
                response._content = body.encode()
                response.headers["Content-Type"] = "text/html"
                with patch.object(self.client.session, "request", return_value=response) as request:
                    with self.assertRaisesRegex(importer.AzureDevOpsError, "AZURE_DEVOPS_PAT"):
                        self.client.list_test_plans()
                self.assertFalse(request.call_args.kwargs["allow_redirects"])

    def test_request_accepts_json_and_empty_responses(self):
        for status, body, expected in ((200, b'{"value": []}', {"value": []}), (204, b"", None)):
            with self.subTest(status=status):
                response = importer.requests.Response()
                response.status_code = status
                response._content = body
                with patch.object(self.client.session, "request", return_value=response):
                    self.assertEqual(self.client.request("GET", "/_apis/testplan/plans"), expected)


class CommandLineTests(unittest.TestCase):
    def test_invalid_excel_stops_before_azure_changes(self):
        with patch.dict(os.environ, {"AZURE_DEVOPS_PAT": "test-token"}), patch("sys.argv", ["importer", str(WORKBOOK)]):
            with patch.object(importer, "read_excel", side_effect=ValueError("Duplicate TC_ID")):
                with patch.object(importer, "AzureDevOpsClient") as client, redirect_stderr(io.StringIO()):
                    self.assertEqual(importer.main(), 2)
                    client.assert_not_called()

    def test_dry_run_needs_no_token_or_network(self):
        output = io.StringIO()
        with patch.dict(os.environ, {}, clear=True), patch("sys.argv", ["importer", str(WORKBOOK), "--dry-run"]):
            with patch.object(importer, "AzureDevOpsClient") as client, redirect_stdout(output):
                self.assertEqual(importer.main(), 0)
                client.assert_not_called()
        self.assertIn("Read 14 test case(s)", output.getvalue())

    def test_live_import_requires_environment_token(self):
        error = io.StringIO()
        with patch.dict(os.environ, {}, clear=True), patch("sys.argv", ["importer", str(WORKBOOK)]):
            with patch.object(importer, "AzureDevOpsClient") as client, redirect_stderr(error):
                self.assertEqual(importer.main(), 2)
                client.assert_not_called()
        self.assertIn("Set AZURE_DEVOPS_PAT", error.getvalue())

    def test_plan_request_failure_stops_import_cleanly(self):
        error = io.StringIO()
        with patch.dict(os.environ, {"AZURE_DEVOPS_PAT": "test-token"}), patch("sys.argv", ["importer", str(WORKBOOK)]):
            with patch.object(importer, "AzureDevOpsClient") as client_type:
                client = client_type.return_value
                client.get_or_create_plan.side_effect = importer.AzureDevOpsError("Authentication required")
                with redirect_stdout(io.StringIO()), redirect_stderr(error):
                    self.assertEqual(importer.main(), 1)
                client.get_or_create_suite.assert_not_called()
                client.create_test_case.assert_not_called()
                client.session.close.assert_called_once()
        self.assertEqual(error.getvalue().strip(), "ERROR: Authentication required")

    def test_import_flow_uses_environment_token_and_links_cases(self):
        with patch.dict(os.environ, {"AZURE_DEVOPS_PAT": "test-token"}), patch("sys.argv", ["importer", str(WORKBOOK)]):
            with patch.object(importer, "AzureDevOpsClient") as client_type, patch.object(importer.time, "sleep"):
                client = client_type.return_value
                client.get_or_create_plan.return_value = {"id": 100}
                client.get_or_create_suite.return_value = {"id": 200}
                client.create_test_case.side_effect = range(1, 15)
                with redirect_stdout(io.StringIO()):
                    self.assertEqual(importer.main(), 0)
                client_type.assert_called_once_with(importer.DEFAULT_ORG, importer.DEFAULT_PROJECT, "test-token")
                self.assertEqual(client.create_test_case.call_count, 14)
                client.add_test_cases_to_suite.assert_called_once_with(100, 200, list(range(1, 15)))
                client.session.close.assert_called_once()

    def test_case_failure_returns_nonzero_but_links_successful_cases(self):
        with patch.dict(os.environ, {"AZURE_DEVOPS_PAT": "test-token"}), patch("sys.argv", ["importer", str(WORKBOOK)]):
            with patch.object(importer, "AzureDevOpsClient") as client_type, patch.object(importer.time, "sleep"):
                client = client_type.return_value
                client.get_or_create_plan.return_value = {"id": 100}
                client.get_or_create_suite.return_value = {"id": 200}
                client.create_test_case.side_effect = [importer.AzureDevOpsError("Rejected row"), *range(2, 15)]
                with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()) as error:
                    self.assertEqual(importer.main(), 1)
                self.assertIn("Rejected row", error.getvalue())
                client.add_test_cases_to_suite.assert_called_once_with(100, 200, list(range(2, 15)))
                client.session.close.assert_called_once()


if __name__ == "__main__":
    unittest.main()