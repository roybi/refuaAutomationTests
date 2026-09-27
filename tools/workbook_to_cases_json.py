"""Convert a MediTik test-case workbook into the docs/setup/*_CASES.json format.

Generic on purpose: pass the workbook path, the test-case sheet, the BDD sheet
and the output path, so the same tool serves every feature workbook
(My Requests, My Appointments, Referrals, ...).

Usage:
    python tools/workbook_to_cases_json.py <workbook.xlsx> <out.json> \
        [--cases-sheet "Test cases"] [--bdd-sheet BDD]
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl

READY = "Ready"
# A blank / "None" Clarification Status in a workbook means nothing blocks the
# case, which we normalise to the explicit "Ready" used by the other suites.
_BLANK_STATUS = {"", "none", "n/a", "-"}


def _rows(worksheet, header_row: int = 1):
    """Split a sheet into (header, data rows).

    ``header_row`` is 1-based: some workbooks put a banner/title line above the
    real header, so the header is not always the first row.
    """
    rows = list(worksheet.iter_rows(values_only=True))
    index = header_row - 1
    if not rows or index >= len(rows):
        return [], []
    header = [str(cell).strip() if cell is not None else "" for cell in rows[index]]
    return header, [r for r in rows[index + 1:] if r and r[0]]


def _record(header, row):
    return {
        header[i]: ("" if row[i] is None else str(row[i]))
        for i in range(len(header))
        if header[i]
    }


def convert(
    workbook_path: Path,
    cases_sheet: str,
    bdd_sheet: str | None,
    cases_header_row: int = 1,
    bdd_header_row: int = 1,
) -> list[dict]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)

    header, rows = _rows(workbook[cases_sheet], cases_header_row)
    cases = [_record(header, row) for row in rows]

    bdd_by_id: dict[str, dict] = {}
    if bdd_sheet and bdd_sheet in workbook.sheetnames:
        bdd_header, bdd_rows = _rows(workbook[bdd_sheet], bdd_header_row)
        for row in bdd_rows:
            record = _record(bdd_header, row)
            bdd_by_id[record.get("TC_ID", "")] = record

    for case in cases:
        status = (case.get("Clarification Status") or "").strip()
        if status.lower() in _BLANK_STATUS:
            status = READY
        case["Clarification Status"] = status
        case["BDD"] = bdd_by_id.get(case.get("TC_ID", ""), {})

    return cases


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--cases-sheet", default="Test cases")
    parser.add_argument("--bdd-sheet", default="BDD")
    parser.add_argument(
        "--cases-header-row",
        type=int,
        default=1,
        help="1-based row holding the cases-sheet header (default 1).",
    )
    parser.add_argument(
        "--bdd-header-row",
        type=int,
        default=1,
        help="1-based row holding the BDD-sheet header (default 1).",
    )
    args = parser.parse_args()

    cases = convert(
        args.workbook,
        args.cases_sheet,
        args.bdd_sheet,
        args.cases_header_row,
        args.bdd_header_row,
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(cases, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    ready = [c["TC_ID"] for c in cases if c["Clarification Status"] == READY]
    print(f"wrote {len(cases)} cases -> {args.output}")
    print(f"ready ({len(ready)}): {', '.join(ready)}")
    print(f"pending: {len(cases) - len(ready)}")


if __name__ == "__main__":
    main()
